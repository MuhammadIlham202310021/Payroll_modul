from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill('solid', fgColor='212529')
HEADER_FONT = Font(bold=True, color='FFFFFF')
DEPT_FILL = PatternFill('solid', fgColor='E9ECEF')
THIN = Side(style='thin', color='BBBBBB')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _header_row(ws, row, headers):
    for col, text in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=text)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = BORDER

# mengisi sheet data: judul, filter, tanggal cetak, nama user, baris header, lalu satu baris per pekerja.
def _sheet_biodata(ws, qs, label, username):
    headers = ['No', 'EMP Code', 'EMP Name', 'ID Card', 'Sex', 'Birth Place',
               'Birth Date', 'Company Begin', 'Department', 'Division']
    widths = [6, 14, 28, 22, 10, 20, 14, 16, 22, 26]
    now = timezone.localtime()

    ws['A1'] = 'REPORT BIODATA'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = f'Filter: {label}'
    ws['A3'] = f'Tanggal cetak: {now:%d-%m-%Y %H:%M}'
    ws['A4'] = f'Dicetak oleh: {username}'

    header_row = 6
    _header_row(ws, header_row, headers)

    row = header_row
    for no, p in enumerate(qs, start=1):
        row += 1
        values = [
            no, p.emp_code, p.emp_name, p.id_card, p.get_sex_display(),
            p.birth_place, p.birth_date, p.company_begin,
            p.division.department.department_name, p.division.division_name,
        ]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=v)
            c.border = BORDER
        ws.cell(row=row, column=7).number_format = 'DD-MM-YYYY'
        ws.cell(row=row, column=8).number_format = 'DD-MM-YYYY'
        ws.cell(row=row, column=1).alignment = Alignment(horizontal='center')

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # mengunci baris header agar tetap terlihat saat menggulir
    ws.freeze_panes = f'A{header_row + 1}'
    # memasang tombol filter di tiap kolom.
    ws.auto_filter.ref = f'A{header_row}:{get_column_letter(len(headers))}{row}'

# memakai data build_summary yang sama dengan halaman Report: baris Department tebal dan abu-abu, baris Division di bawahnya, dan baris TOTAL di akhir.
def _sheet_ringkasan(ws, summary):
    headers = ['Department Code', 'Department Name', 'Division Code', 'Division Name', 'Jumlah']
    widths = [18, 26, 16, 28, 10]

    ws['A1'] = 'RINGKASAN JUMLAH PEKERJA'
    ws['A1'].font = Font(bold=True, size=14)

    _header_row(ws, 3, headers)

    row = 3
    grand_total = 0
    for d in summary:
        row += 1
        for col, v in enumerate([d['code'], d['name'], '', '', d['total']], start=1):
            c = ws.cell(row=row, column=col, value=v)
            c.font = Font(bold=True)
            c.fill = DEPT_FILL
            c.border = BORDER
        grand_total += d['total']

        for v in d['divisions']:
            row += 1
            for col, val in enumerate([d['code'], d['name'], v['code'], v['name'], v['total']], start=1):
                ws.cell(row=row, column=col, value=val).border = BORDER

    row += 1
    ws.cell(row=row, column=1, value='TOTAL').font = Font(bold=True)
    ws.cell(row=row, column=5, value=grand_total).font = Font(bold=True)
    for col in range(1, 6):
        ws.cell(row=row, column=col).border = BORDER

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_workbook(qs, summary, label, username):
    wb = Workbook()
    ws1 = wb.active
    ws1.title = 'Biodata'
    _sheet_biodata(ws1, qs, label, username)

    ws2 = wb.create_sheet('Ringkasan')
    _sheet_ringkasan(ws2, summary)
    return wb